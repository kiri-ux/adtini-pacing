"""XLSX exports laid out the way the buying team's sheets already are.

The point is not a data dump - it is that a buyer can open the export and see
the same columns in the same order as the tab they keep by hand, so the
handover costs nothing.
"""
from __future__ import annotations

import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from models import PACING_CLICK, PACING_EVENT, PACING_IMPRESSION

NAVY = "FF123A63"
FILL = "FFF1F2F4"
HEAD = Font(bold=True, color="FFFFFFFF", size=10)
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
TITLE = Font(bold=True, size=12)

# Column headers per pacing type, matching the three tabs.
HEADERS = {
    PACING_IMPRESSION: [
        "Campaign Elements:", "Impr.", "Tot. Impr.", "Day Impr.", "CPM",
        "Day Budg.", "Tot Budg.", "TD Impr.", "Impr. Left", "On Pace", "Pacing",
    ],
    PACING_CLICK: [
        "Campaign Elements", "Mon Spend", "Total Spend", "Daily Spend", "CPC",
        "Mon Clicks", "Total Clicks", "Daily Clicks", "TD Spend", "Spend Left",
        "On Pace", "Pacing",
    ],
    PACING_EVENT: [
        "Campaign Elements", "Client Monthly Budget", "Client Total Budget",
        "Google Monthly Spend", "Google Total Spend", "Google Daily Spend",
        "Google CPE", "Monthly Events", "Total Events", "Daily Events",
        "TD Spend", "Spend Left", "On Pace", "Pacing",
    ],
}

MONEY = "$#,##0.00"
COUNT = "#,##0;(#,##0)"
PCT = "0.00%;(0.00%)"


def _row_values(row, pacing_type: str, line_item=None) -> list:
    if pacing_type == PACING_IMPRESSION:
        return [
            row.label, row.monthly_target, row.total_target, row.daily_target,
            row.unit_cost, row.daily_budget, row.total_budget, row.to_date,
            row.remaining, row.on_pace, row.pacing_pct,
        ]
    if pacing_type == PACING_CLICK:
        clicks_total = (row.total_target / row.unit_cost) if row.unit_cost else None
        clicks_month = (row.monthly_target / row.unit_cost) if row.unit_cost else None
        clicks_daily = (
            clicks_total / row.flight_days if clicks_total and row.flight_days else None
        )
        return [
            row.label, row.monthly_target, row.total_target, row.daily_target,
            row.unit_cost, clicks_month, clicks_total, clicks_daily,
            row.to_date, row.remaining, row.on_pace, row.pacing_pct,
        ]
    li = line_item
    return [
        row.label,
        getattr(li, "client_monthly_budget", None),
        getattr(li, "client_total_budget", None),
        row.monthly_target, row.total_target, row.daily_target, row.unit_cost,
        getattr(li, "monthly_events", None), getattr(li, "total_events", None),
        (
            (li.total_events / row.flight_days)
            if li is not None and li.total_events and row.flight_days
            else None
        ),
        row.to_date, row.remaining, row.on_pace, row.pacing_pct,
    ]


def _formats(pacing_type: str) -> list[str | None]:
    if pacing_type == PACING_IMPRESSION:
        return [None, COUNT, COUNT, "#,##0.00", MONEY, MONEY, MONEY, COUNT, COUNT, COUNT, PCT]
    if pacing_type == PACING_CLICK:
        return [None, MONEY, MONEY, MONEY, MONEY, COUNT, COUNT, COUNT, MONEY, MONEY, MONEY, PCT]
    return [
        None, MONEY, MONEY, MONEY, MONEY, MONEY, MONEY, COUNT, COUNT, COUNT,
        MONEY, MONEY, MONEY, PCT,
    ]


def order_workbook(view) -> Workbook:
    """One order, laid out as its section in the pacing sheet."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Pacing"

    pacing_type = view.order.pacing_type
    headers = HEADERS[pacing_type]
    formats = _formats(pacing_type)

    sheet.append([view.client.name, "Run Dates:", view.order.start_date,
                  view.order.end_date, "", "Days", view.total.flight_days])
    sheet["A1"].font = TITLE
    sheet.append([])

    sheet.append(headers + [d.strftime("%-d-%b") for d in view.grid_dates])
    for cell in sheet[3]:
        cell.font = HEAD
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="left")
    sheet.append([])

    by_id = {li.id: li for li in view.order.line_items}
    for row in view.rows:
        values = _row_values(row, pacing_type, by_id.get(row.line_item_id))
        daily = view.grid.get(row.line_item_id, {})
        sheet.append(values + [daily.get(d) for d in view.grid_dates])
        for idx, fmt in enumerate(formats, start=1):
            if fmt:
                sheet.cell(row=sheet.max_row, column=idx).number_format = fmt
        for idx in range(len(headers) + 1, len(headers) + len(view.grid_dates) + 1):
            sheet.cell(row=sheet.max_row, column=idx).number_format = COUNT

    sheet.append([])
    total_values = _row_values(view.total, pacing_type)
    total_values[0] = "Total:"
    totals_by_day = [
        sum(
            (view.grid.get(r.line_item_id, {}).get(d) or 0) for r in view.rows
        )
        for d in view.grid_dates
    ]
    sheet.append(total_values + totals_by_day)
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    for idx, fmt in enumerate(formats, start=1):
        if fmt:
            sheet.cell(row=sheet.max_row, column=idx).number_format = fmt

    # The flat daily target, the way the sheet carries it under the total.
    sheet.append(
        [""] * len(headers) + [view.on_pace_daily] * len(view.grid_dates)
    )

    sheet.column_dimensions["A"].width = 34
    for idx in range(2, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(idx)].width = 14
    sheet.freeze_panes = f"B{4}"
    return book


OVERVIEW_HEADERS = [
    "Buyer", "Market", "Client Name", "Order", "Type", "Start Date:", "End Date:",
    "Sold Total", "Mon. Sold", "CTR", "Total Delivered", "Total On Pace",
    "Total Pacing:", "TD Mon.", "Mon. On Pace", "Mon. Pacing", "Adjusted on:",
    "Monthly Notes:",
]


def overview_workbook(rows) -> Workbook:
    """Every order on one line - the summary tab."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Overview"

    sheet.append(OVERVIEW_HEADERS)
    for cell in sheet[1]:
        cell.font = HEAD
        cell.fill = HEAD_FILL

    for item in rows:
        total = item.total
        sheet.append([
            item.order.buyer,
            item.client.market,
            item.client.name,
            item.order.name,
            item.order.pacing_type,
            item.order.start_date,
            item.order.end_date,
            total.total_target,
            total.monthly_target,
            total.ctr,
            total.to_date,
            total.on_pace,
            total.pacing_delta,
            total.month_to_date,
            total.month_on_pace,
            total.month_pacing_pct,
            item.order.last_adjusted_on,
            item.order.adjustment_note or item.order.notes,
        ])
        current = sheet.max_row
        for col in (8, 9, 11, 12, 13, 14, 15):
            sheet.cell(row=current, column=col).number_format = COUNT
        sheet.cell(row=current, column=10).number_format = "0.00%"
        sheet.cell(row=current, column=16).number_format = PCT

    widths = [12, 26, 30, 30, 12, 12, 12, 14, 14, 9, 15, 15, 14, 14, 14, 12, 13, 26]
    for idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    sheet.freeze_panes = "D2"
    return book
